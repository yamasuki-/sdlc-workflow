"""ワークフロー本体の自己テスト。  実行: python3 -m unittest discover -s tests -v"""
from __future__ import annotations

import contextlib
import importlib.machinery
import io
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import textwrap
import unittest
from argparse import Namespace
from pathlib import Path

# サンプルの src を同じ秒・同じサイズで書き換えるため、.pyc の再利用で古いコードが走らないようにする
os.environ["PYTHONDONTWRITEBYTECODE"] = "1"

# ~/.claude をテスト用の一時ディレクトリにする（実環境の ~/.claude に触れない）
TEST_HOME = Path(tempfile.mkdtemp(prefix="sdlc-home-"))
os.environ["HOME"] = str(TEST_HOME)

REPO = Path(__file__).resolve().parents[1]
BIN = REPO / ".claude" / "sdlc" / "bin"
HOOKS = REPO / ".claude" / "hooks" / "sdlc"
INSTALLER = REPO / "install.py"
sys.path.insert(0, str(REPO / ".claude" / "sdlc" / "lib"))

from sdlc_lib import checks, defects, fm, gates, records, state, tasks, tdd, trace  # noqa: E402
from sdlc_lib.config import load_project, phase_of_path, project_problems  # noqa: E402


def load_script(name):
    loader = importlib.machinery.SourceFileLoader(name, str(BIN / name))
    spec = importlib.util.spec_from_loader(name, loader)
    mod = importlib.util.module_from_spec(spec)
    loader.exec_module(mod)
    return mod


approve = load_script("approve")

PROJECT_YAML = """\
workflow: sdlc
language: python
test_framework: unittest
commands:
  test: "python3 -m unittest discover -s tests -t . -q"
src_dir: src
test_dir: tests
"""


class Repo:
    """一時 git リポジトリ上のサンプルプロジェクト。"""

    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="sdlc-test-"))
        self.run("git", "init", "-q")
        self.run("git", "config", "user.email", "t@example.com")
        self.run("git", "config", "user.name", "tester")
        self.write("project.yaml", PROJECT_YAML)
        self.write(".gitignore", "__pycache__/\n.sdlc/\n")
        for d in ("src", "tests", "tests/unit", "tests/regression", "tests/integration", "tests/system"):
            self.write(f"{d}/__init__.py", "")
        self.write("input/req.md", "# 要求\n\n- 足し算ができること\n")
        state.save(self.dir, dict(state.DEFAULT_STATE))

    def run(self, *cmd, **kw):
        return subprocess.run(cmd, cwd=self.dir, capture_output=True, text=True, **kw)

    def write(self, rel, text):
        p = self.dir / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")

    @property
    def proj(self):
        return load_project(self.dir)

    def approve(self, phase, decision="approve"):
        a = Namespace(target=phase, defect_id=None, decision=decision, approver="tester", comment="ok", yes=True,
                      option=None)
        with contextlib.redirect_stdout(io.StringIO()):
            approve.approve_phase(self.dir, self.proj, a)

    def hook(self, script, tool_input, env=None):
        data = {"tool_input": tool_input, "cwd": str(self.dir)}
        e = dict(os.environ, CLAUDE_PROJECT_DIR=str(self.dir))
        e.pop("SDLC_MAINTENANCE", None)
        e.update(env or {})
        return subprocess.run([sys.executable, str(HOOKS / script)], input=json.dumps(data), capture_output=True,
                              text=True, cwd=self.dir, env=e)

    def can_write(self, rel, **kw):
        r = self.hook("guard_write.py", {"file_path": str(self.dir / rel), "content": "x"}, **kw)
        return r.returncode == 0, r.stderr

    def cleanup(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class Base(unittest.TestCase):
    def setUp(self):
        self.r = Repo()
        self._cwd = os.getcwd()
        os.chdir(self.r.dir)

    def tearDown(self):
        os.chdir(self._cwd)
        self.r.cleanup()

    def approve_upto(self, k):
        """工程 0..k の成果物を用意して承認する。"""
        docs = {
            0: ("docs/00-intake/req.md", "### SRC-001\n\n> 足し算\n"),
            1: ("docs/01-requirements/requirements.md", "### REQ-001: 足し算\n- 追跡元: SRC-001\n"),
            2: ("docs/02-system-requirements/sr.md",
                "### SR-001: add\n- 追跡元: REQ-001\n\n### AT-001: add の受入\n- 追跡元: SR-001\n"),
            3: ("docs/03-basic-design/bd.md", "### BD-001: calc\n- 追跡元: SR-001\n"),
            4: ("docs/04-detailed-design/dd.md",
                "### DD-001: add()\n- 追跡元: BD-001\n\n### UT-001: add のテスト\n- 追跡元: DD-001\n"),
        }
        for i in range(k + 1):
            if i in docs:
                self.r.write(*docs[i])
            if i == 4:
                self.r.write("tests/unit/test_calc.py", textwrap.dedent("""\
                    # trace: UT-001
                    import unittest
                    from src.calc import add

                    class T(unittest.TestCase):
                        def test_add(self):
                            self.assertEqual(add(1, 2), 3)
                    """))
            self.r.approve(i)


class TestConfig(Base):
    def test_project_and_areas(self):
        p = self.r.proj
        self.assertEqual(project_problems(p), [])
        self.assertEqual(p["test_dirs"]["unit"], "tests/unit")
        self.assertEqual(phase_of_path("src/a.py", p), "05-implementation")
        self.assertEqual(phase_of_path("tests/unit/t.py", p), "04-detailed-design")
        self.assertEqual(phase_of_path("tests/regression/t.py", p), None)
        self.assertEqual(phase_of_path("input/x.xlsx", p), "00-intake")

    def test_reserved_dirs(self):
        self.r.write("project.yaml", PROJECT_YAML.replace("src_dir: src", "src_dir: docs/src"))
        self.assertTrue(any("予約" in x for x in project_problems(self.r.proj)))


class TestGates(Base):
    def test_hash_invalidation(self):
        self.approve_upto(0)
        self.assertTrue(gates.gate_status(self.r.dir, self.r.proj, 0)["valid"])
        self.r.write("docs/00-intake/req.md", "changed")
        st = gates.gate_status(self.r.dir, self.r.proj, 0)
        self.assertFalse(st["valid"])
        self.assertIn("ハッシュ不一致", st["reason"])

    def test_cannot_approve_without_prereq(self):
        self.r.write("docs/01-requirements/r.md", "x")
        with self.assertRaises(SystemExit):
            self.r.approve(1)

    def test_rejected_not_valid(self):
        self.r.write("docs/00-intake/req.md", "x")
        self.r.approve(0, "reject")
        self.assertFalse(gates.gate_status(self.r.dir, self.r.proj, 0)["valid"])

    def test_history_kept(self):
        self.approve_upto(0)
        self.r.write("docs/00-intake/req.md", "v2")
        self.r.approve(0)
        _, body = fm.read(gates.gate_path(self.r.dir, 0))
        self.assertIn("## 履歴", body)


class TestGuardWrite(Base):
    def test_gates_dir_denied(self):
        ok, msg = self.r.can_write("state/gates/00-intake.md")
        self.assertFalse(ok)
        ok, _ = self.r.can_write("state/gates/x.md", env={"SDLC_MAINTENANCE": "1"})
        self.assertFalse(ok)

    def test_protected_paths(self):
        self.assertFalse(self.r.can_write(".claude/settings.json")[0])
        self.assertFalse(self.r.can_write(".claude/sdlc/lib/sdlc_lib/gates.py")[0])
        self.assertFalse(self.r.can_write(".claude/hooks/sdlc/guard_write.py")[0])
        self.assertFalse(self.r.can_write("scripts/sdlc")[0])
        self.assertFalse(self.r.can_write(".claude/skills/phase-intake/SKILL.md")[0])
        self.assertTrue(self.r.can_write(".claude/skills/my-own-skill/SKILL.md")[0])
        self.assertTrue(self.r.can_write(".claude/settings.json", env={"SDLC_MAINTENANCE": "1"})[0])

    def test_phase_order(self):
        self.assertTrue(self.r.can_write("docs/00-intake/a.md")[0])
        ok, msg = self.r.can_write("docs/01-requirements/a.md")
        self.assertFalse(ok)
        self.assertIn("00-intake", msg)
        self.approve_upto(0)
        self.assertTrue(self.r.can_write("docs/01-requirements/a.md")[0])
        # 承認済み工程は変更不可
        ok, msg = self.r.can_write("docs/00-intake/a.md")
        self.assertFalse(ok)
        self.assertIn("承認済み", msg)

    def test_free_paths(self):
        self.assertTrue(self.r.can_write("state/tasks.yaml")[0])
        self.assertTrue(self.r.can_write("tests/regression/test_x.py")[0])
        self.assertTrue(self.r.can_write("README.md")[0])

    def test_tdd(self):
        self.approve_upto(4)  # 単体テストは Red（src/calc.py が無い）
        ok, msg = self.r.can_write("src/calc.py")
        self.assertTrue(ok, msg)
        self.r.write("src/calc.py", "def add(a, b):\n    return a + b\n")
        ok, msg = self.r.can_write("src/calc.py")
        self.assertTrue(ok, msg)  # Red 観測済み → Refactor 許可
        tdd.reset(self.r.dir)
        ok, msg = self.r.can_write("src/calc.py")
        self.assertFalse(ok)
        self.assertIn("TDD", msg)

    def test_unit_tests_frozen_after_detailed_design(self):
        self.approve_upto(4)
        self.assertFalse(self.r.can_write("tests/unit/test_calc.py")[0])


class TestGuardBashAndSkill(Base):
    def bash(self, cmd):
        return self.r.hook("guard_bash.py", {"command": cmd}).returncode == 0

    def test_bash(self):
        self.assertTrue(self.bash("cat state/gates/00-intake.md"))
        self.assertFalse(self.bash("echo x > state/gates/00-intake.md"))
        self.assertFalse(self.bash("scripts/approve 0 --yes"))
        self.assertFalse(self.bash("python3 scripts/approve 0"))
        self.assertFalse(self.bash("cd /tmp && ./scripts/approve defect DEF-001"))
        self.assertTrue(self.bash("git commit -m 'gate は scripts/approve で記録'"))
        self.assertTrue(self.bash("grep -n approve scripts/approve"))
        self.assertFalse(self.bash("echo '{}' > .claude/settings.local.json"))
        self.assertFalse(self.bash("echo x > docs/01-requirements/a.md"))
        self.assertFalse(self.bash("sed -i s/a/b/ src/calc.py"))
        self.assertFalse(self.bash("cp /tmp/x .claude/hooks/sdlc/guard_write.py"))
        self.assertFalse(self.bash("python3 install.py --user"))
        self.assertTrue(self.bash("python3 tools/install.py"))
        self.assertFalse(self.bash("echo x > ~/.claude/sdlc/RULES.md"))
        self.assertTrue(self.bash("echo x > ~/.claude/skills/my-own-skill/notes.md"))
        self.assertTrue(self.bash("python3 -m unittest discover -s tests 2>&1 | tail"))
        self.assertTrue(self.bash("scripts/sdlc status"))
        self.assertTrue(self.bash("git merge --no-ff sdlc/T-001"))
        self.assertTrue(self.bash("ls docs > /dev/null"))

    def test_skill(self):
        r = self.r.hook("guard_skill.py", {"skill": "phase-requirements"})
        self.assertEqual(r.returncode, 2)
        self.approve_upto(0)
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "phase-requirements"}).returncode, 0)
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "sdlc"}).returncode, 0)


class TestSettingsProtection(Base):
    def test_settings_in_and_out_of_repo(self):
        # SDLC プロジェクトの設定ファイル
        self.assertFalse(self.r.can_write(".claude/settings.local.json")[0])
        # ~/.claude の設定ファイルはどこからでも保護（メンテナンスモードでのみ可）
        user_settings = TEST_HOME / ".claude" / "settings.json"
        r = self.r.hook("guard_write.py", {"file_path": str(user_settings), "content": "{}"})
        self.assertEqual(r.returncode, 2)
        r = self.r.hook("guard_write.py", {"file_path": str(user_settings), "content": "{}"},
                        env={"SDLC_MAINTENANCE": "1"})
        self.assertEqual(r.returncode, 0)
        # 関係のない場所の .claude/settings.json は制限しない
        outside = Path(tempfile.mkdtemp(prefix="sdlc-other-"))
        try:
            r = self.r.hook("guard_write.py", {"file_path": str(outside / ".claude" / "settings.json"), "content": "{}"})
            self.assertEqual(r.returncode, 0)
        finally:
            shutil.rmtree(outside, ignore_errors=True)


class TestChecks(Base):
    def set_checks(self, block):
        self.r.write("project.yaml", PROJECT_YAML + block)

    def test_levels_and_skip(self):
        self.set_checks('checks:\n  common:\n    must: ["true"]\n    should: ["false"]\n'
                        '    may: ["no_such_tool_xyz --version"]\n')
        res = checks.run(self.r.dir, self.r.proj, 0)
        self.assertTrue(res["must_passed"])
        self.assertEqual(res["counts"]["should"]["fail"], 1)
        self.assertEqual(res["skipped"], ["no_such_tool_xyz"])
        self.assertIn("SHOULD", checks.render(res))

    def test_must_failure_blocks_approval(self):
        self.set_checks('checks:\n  00-intake:\n    must: [{name: lint, cmd: "false"}]\n')
        self.r.write("docs/00-intake/req.md", "### SRC-001\n")
        with self.assertRaises(SystemExit):
            self.r.approve(0)
        self.r.approve(0, "reject")  # 差戻しは記録できる
        self.assertEqual(fm.read(gates.gate_path(self.r.dir, 0))[0]["decision"], "rejected")

    def test_failed_review_blocks_approval(self):
        self.r.write("docs/00-intake/req.md", "### SRC-001\n")
        self.r.write("docs/00-intake/review.md", "# レビュー\n\n- 判定: 不合格\n- 残指摘: 1\n")
        with self.assertRaises(SystemExit):
            self.r.approve(0)

    def test_builtin_trace_check(self):
        self.r.write("docs/01-requirements/r.md", "### REQ-001: x\n- 追跡元: SRC-404\n")
        res = checks.run(self.r.dir, self.r.proj, 1)
        self.assertFalse(res["must_passed"])


class TestRecords(Base):
    def test_decisions_and_questions(self):
        self.assertEqual(records.add_decision(self.r.dir, "言語は Python", "00-intake", "人の指示"), "DEC-001")
        self.assertEqual(records.add_decision(self.r.dir, "DB は SQLite"), "DEC-002")
        q1 = records.add_question(self.r.dir, "同時貸出の上限は？", "高", "01-requirements")
        q2 = records.add_question(self.r.dir, "期限の起算日は？")
        self.assertEqual((q1, q2), ("OQ-001", "OQ-002"))
        records.answer_question(self.r.dir, q1, "3 冊")
        open_ids = [q["id"] for q in records.questions(self.r.dir, only_open=True)]
        self.assertEqual(open_ids, ["OQ-002"])
        self.assertIn("回答: 3 冊", records.questions_path(self.r.dir).read_text(encoding="utf-8"))


class TestTasks(unittest.TestCase):
    def data(self):
        return {"tasks": [
            {"id": "T1", "agent": "analyst", "files": ["docs/a.md"]},
            {"id": "T2", "agent": "analyst", "files": ["docs/a.md"]},
            {"id": "T3", "agent": "analyst", "files": ["docs/b.md"]},
            {"id": "T4", "agent": "analyst", "files": ["docs/c.md"], "depends_on": ["T1"]},
        ]}

    def test_ready_respects_files_and_deps(self):
        d = self.data()
        for t in d["tasks"]:
            t.setdefault("status", "pending")
            t.setdefault("depends_on", [])
        self.assertEqual([t["id"] for t in tasks.ready(d, 2)], ["T1", "T3"])
        tasks.set_status(d, "T1", "in_progress")
        self.assertEqual([t["id"] for t in tasks.ready(d, 2)], ["T3"])
        tasks.set_status(d, "T1", "done")
        self.assertEqual([t["id"] for t in tasks.ready(d, 2)], ["T2", "T3"])

    def test_cycle(self):
        d = {"tasks": [{"id": "A", "agent": "analyst", "files": ["x"], "depends_on": ["B"], "status": "pending"},
                       {"id": "B", "agent": "analyst", "files": ["y"], "depends_on": ["A"], "status": "pending"}]}
        self.assertTrue(any("循環" in e for e in tasks.validate(d)))

    def test_recover(self):
        d = {"tasks": [{"id": "A", "status": "in_progress"}]}
        self.assertEqual(tasks.recover(d), ["A"])
        self.assertEqual(d["tasks"][0]["status"], "pending")


class TestTrace(Base):
    def test_check_and_impact(self):
        self.approve_upto(4)
        g = trace.build(self.r.dir, self.r.proj)
        self.assertEqual(trace.check(g, 4), [])
        issues = trace.check(g, 5)
        self.assertTrue(any("DD-001" in i and "コード" in i for i in issues))
        imp = trace.impact(g, ["SR-001"])
        self.assertIn("UT-001", imp["ids"])
        self.assertIn("tests/unit/test_calc.py", imp["files"])
        self.assertIn("REQ-001", trace.upstream(g, ["DD-001"]))

    def test_exclude_and_manual(self):
        self.r.write("docs/00-intake/a.md", "### SRC-001\n\n### SRC-002\n")
        self.r.write("docs/01-requirements/r.md", "### REQ-001: x\n- 追跡元: SRC-001\n\n## 対象外\n- 除外: SRC-002（表題）\n")
        g = trace.build(self.r.dir, self.r.proj)
        self.assertEqual([i for i in trace.check(g, 1)], [])
        self.r.write("docs/02-system-requirements/at.md", "### AT-001: y\n- 追跡元: REQ-001\n")
        self.r.write("docs/07-system-test/manual-tests.md", "## 記録1\n- trace: AT-001\n- 結果: 合格\n")
        g = trace.build(self.r.dir, self.r.proj)
        self.assertFalse(any("AT-001" in i for i in trace.check(g, 7)))

    def test_dangling(self):
        self.r.write("docs/01-requirements/r.md", "### REQ-002: x\n- 追跡元: SRC-999\n")
        issues = trace.check(trace.build(self.r.dir, self.r.proj), 1)
        self.assertTrue(any("未定義" in i for i in issues))


class TestDefectFlow(Base):
    def sdlc(self, *args):
        r = subprocess.run([sys.executable, str(BIN / "sdlc"), *args], cwd=self.r.dir, capture_output=True, text=True,
                           env=dict(os.environ, CLAUDE_PROJECT_DIR=str(self.r.dir)))
        return r

    def test_full_defect_cycle(self):
        self.approve_upto(4)
        self.r.write("src/calc.py", "def add(a, b):\n    # trace: DD-001\n    return a - b\n")
        self.r.write("tests/unit/test_calc.py", (self.r.dir / "tests/unit/test_calc.py").read_text())
        st = state.load(self.r.dir)
        st["current_phase"] = "06-integration-test"
        state.save(self.r.dir, st)
        # 起票（再現テストは regression に）
        r = self.sdlc("defect", "new", "--title", "add が引き算", "--detected-in", "06", "--by", "auto",
                      "--test", "tests/regression/test_def_001.py")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.r.write("tests/regression/test_def_001.py",
                     "import unittest\nfrom src.calc import add\nclass T(unittest.TestCase):\n"
                     "    def test(self):\n        self.assertEqual(add(2, 2), 4)\n")
        meta, body = defects.load(self.r.dir, "DEF-001")
        meta.update(cause_ids=["DD-001"], injected_phase="05-implementation", escaped_phase="05-implementation")
        defects.save(self.r.dir, "DEF-001", meta, body)
        r = self.sdlc("defect", "impact", "DEF-001", "--write")
        self.assertEqual(r.returncode, 0, r.stderr)
        meta, _ = defects.load(self.r.dir, "DEF-001")
        self.assertIn("05-implementation", meta["rework_phases"])
        # 実装工程の承認（コードが存在する状態）
        self.r.approve(5)
        self.assertFalse(self.r.can_write("src/calc.py")[0])
        self.assertEqual(self.sdlc("defect", "analyzed", "DEF-001").returncode, 0)
        self.assertNotEqual(self.sdlc("defect", "start", "DEF-001").returncode, 0)  # 未承認
        with contextlib.redirect_stdout(io.StringIO()):
            approve.approve_defect(self.r.dir, self.r.proj, Namespace(
                target="defect", defect_id="DEF-001", decision="approve", approver="tester", comment="", yes=True,
                option=None))
        self.assertIn("再承認要", gates.gate_status(self.r.dir, self.r.proj, 5)["reason"])
        r = self.sdlc("defect", "start", "DEF-001")
        self.assertEqual(r.returncode, 0, r.stderr)
        tdd.reset(self.r.dir)
        ok, msg = self.r.can_write("src/calc.py")
        self.assertTrue(ok, msg)  # 再現テストが Red
        self.r.write("src/calc.py", "def add(a, b):\n    # trace: DD-001\n    return a + b\n")
        self.assertNotEqual(self.sdlc("defect", "close", "DEF-001").returncode, 0)  # 工程5未再承認
        self.r.approve(5)
        self.assertEqual(self.sdlc("phase", "advance").returncode, 0)
        r = self.sdlc("defect", "close", "DEF-001")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        st = state.load(self.r.dir)
        self.assertEqual((st["mode"], st["current_phase"]), ("normal", "06-integration-test"))

    def test_status_next_action(self):
        r = self.sdlc("status", "--json")
        self.assertEqual(json.loads(r.stdout)["next"]["action"], "run_phase")
        self.approve_upto(0)
        self.assertEqual(json.loads(self.sdlc("status", "--json").stdout)["next"]["action"], "advance")
        self.assertEqual(self.sdlc("phase", "advance").returncode, 0)
        self.sdlc("phase", "step", "gate")
        self.assertEqual(json.loads(self.sdlc("status", "--json").stdout)["next"]["action"], "await_approval")


class TestTestRun(Base):
    def sdlc(self, *args):
        return subprocess.run([sys.executable, str(BIN / "sdlc"), *args], cwd=self.r.dir, capture_output=True,
                              text=True, env=dict(os.environ, CLAUDE_PROJECT_DIR=str(self.r.dir)))

    def test_red_then_green(self):
        self.approve_upto(3)
        self.r.write("docs/04-detailed-design/dd.md",
                     "### DD-001: add()\n- 追跡元: BD-001\n\n### UT-001: t\n- 追跡元: DD-001\n")
        self.r.write("tests/unit/test_calc.py", "# trace: UT-001\nimport unittest\nfrom src.calc import add\n"
                     "class T(unittest.TestCase):\n    def test(self):\n        self.assertEqual(add(1, 2), 3)\n")
        st = state.load(self.r.dir)
        st["current_phase"] = "04-detailed-design"
        state.save(self.r.dir, st)
        r = self.sdlc("test-run", "--expect", "red")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertTrue((self.r.dir / "docs/04-detailed-design/test-run-red.md").exists())
        self.assertEqual(self.sdlc("test-run", "--expect", "green").returncode, 5)


class TestChangeFlow(Base):
    def sdlc(self, *args):
        return subprocess.run([sys.executable, str(BIN / "sdlc"), *args], cwd=self.r.dir, capture_output=True,
                              text=True, env=dict(os.environ, CLAUDE_PROJECT_DIR=str(self.r.dir)))

    def test_change_request_cycle(self):
        self.approve_upto(2)
        self.r.write("tests/regression/test_ok.py",
                     "import unittest\nclass T(unittest.TestCase):\n    def test(self):\n        pass\n")
        st = state.load(self.r.dir)
        st["current_phase"], st["step"] = "03-basic-design", "execute"
        state.save(self.r.dir, st)
        r = self.sdlc("change", "new", "--title", "引き算も必要")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue((self.r.dir / "state/changes/CR-001.md").exists())
        meta, body = defects.load(self.r.dir, "CR-001")
        meta.update(cause_ids=["REQ-001"], start_phase="01-requirements",
                    options=[{"id": "A", "type": "ideal", "title": "演算を汎用化"},
                             {"id": "B", "type": "minimal", "title": "sub を追加"}], recommended="A")
        defects.save(self.r.dir, "CR-001", meta, body)
        r = self.sdlc("change", "impact", "CR-001", "--write")
        self.assertEqual(r.returncode, 0, r.stderr)
        meta, _ = defects.load(self.r.dir, "CR-001")
        self.assertEqual(meta["rework_phases"], ["01-requirements", "02-system-requirements"])
        self.assertEqual(self.sdlc("change", "analyzed", "CR-001").returncode, 0)
        with contextlib.redirect_stdout(io.StringIO()):
            approve.approve_defect(self.r.dir, self.r.proj, Namespace(
                target="change", defect_id="CR-001", decision="approve", approver="tester", comment="", yes=True,
                option="B"))
        self.assertTrue((self.r.dir / "state/gates/change-CR-001.md").exists())
        self.assertEqual(defects.load(self.r.dir, "CR-001")[0]["selected_option"], "B")
        self.assertIn("再承認要", gates.gate_status(self.r.dir, self.r.proj, 1)["reason"])
        r = self.sdlc("change", "start", "CR-001")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(state.load(self.r.dir)["current_phase"], "01-requirements")
        self.assertTrue(self.r.can_write("docs/01-requirements/requirements.md")[0])
        # 再開先工程（03）の影響範囲外の作業は継続できる
        self.assertTrue(self.r.can_write("docs/03-basic-design/other.md")[0])
        self.r.write("docs/01-requirements/requirements.md", "### REQ-001: 足し算と引き算\n- 追跡元: SRC-001\n")
        self.r.approve(1)
        self.assertEqual(self.sdlc("phase", "advance").returncode, 0)
        self.r.approve(2)
        self.assertEqual(self.sdlc("phase", "advance").returncode, 0)
        r = self.sdlc("change", "close", "CR-001")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        st = state.load(self.r.dir)
        self.assertEqual((st["mode"], st["current_phase"], st["step"]), ("normal", "03-basic-design", "execute"))

    def test_status_routes_to_change(self):
        self.sdlc("change", "new", "--title", "x")
        nxt = json.loads(self.sdlc("status", "--json").stdout)["next"]
        self.assertEqual((nxt["action"], nxt["skill"]), ("change", "change"))


class TestAdapters(Base):
    def convert(self, rel):
        from intake import convert
        with contextlib.redirect_stdout(io.StringIO()):
            convert.main([rel])
        return (self.r.dir / "docs/00-intake" / (Path(rel).stem + ".md")).read_text(encoding="utf-8")

    def test_md(self):
        self.r.write("input/spec.md", "# 概要\n\n## 機能\n\n- ログインできる\n- 検索できる\n\n| ID | 内容 |\n|---|---|\n| 1 | 帳票 |\n")
        out = self.convert("input/spec.md")
        self.assertIn("SRC-001", out)
        self.assertIn("概要 > 機能 / L5", out)
        self.assertIn("内容: 帳票", out)

    def test_xlsx_merged(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "要求一覧"
        ws.append(["分類", "ID", "要求"])
        ws.append(["認証", "R1", "ログインできる"])
        ws.append([None, "R2", "ログアウトできる"])
        ws.merge_cells("A2:A3")
        wb.save(self.r.dir / "input/req.xlsx")
        out = self.convert("input/req.xlsx")
        self.assertIn("要求一覧!A3:C3", out)
        self.assertIn("分類: 認証\nID: R2", out.replace("> ", ""))
        self.assertIn("確信度: 中", out)

    def test_docx(self):
        import docx
        d = docx.Document()
        d.add_heading("機能要求", 1)
        d.add_paragraph("利用者は検索できること。")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "項目", "内容"
        t.cell(1, 0).text, t.cell(1, 1).text = "応答", "3秒以内"
        d.save(self.r.dir / "input/req.docx")
        out = self.convert("input/req.docx")
        self.assertIn("機能要求 / 段落2", out)
        self.assertIn("表1 行2", out)
        self.assertIn("内容: 3秒以内", out)

    def test_pptx(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "権限"
        rows = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        rows.cell(0, 0).text, rows.cell(0, 1).text = "機能", "管理者"
        rows.cell(1, 0).text, rows.cell(1, 1).text = "削除", "○"
        prs.save(self.r.dir / "input/deck.pptx")
        out = self.convert("input/deck.pptx")
        self.assertIn("スライド1「権限」", out)
        self.assertIn("管理者: ○", out)
        self.assertIn("確信度: 中", out)

    def test_refuse_after_gate(self):
        self.approve_upto(0)
        with self.assertRaises(SystemExit):
            self.convert("input/req.md")



def install(*args, home=None):
    env = dict(os.environ, HOME=str(home or TEST_HOME))
    return subprocess.run([sys.executable, str(INSTALLER), *args], capture_output=True, text=True, env=env)


def run_hook(hook_path, project, tool_input, env=None):
    data = {"tool_input": tool_input, "cwd": str(project)}
    e = dict(os.environ, CLAUDE_PROJECT_DIR=str(project))
    e.pop("SDLC_MAINTENANCE", None)
    e.update(env or {})
    return subprocess.run([sys.executable, str(hook_path)], input=json.dumps(data), capture_output=True, text=True,
                          cwd=project, env=e)


class TestPlacement(unittest.TestCase):
    """~/.claude への配置・プロジェクトの .claude による上書き・プロジェクト単体での配置。"""

    def setUp(self):
        self.home = Path(tempfile.mkdtemp(prefix="sdlc-home-"))
        self.user = self.home / ".claude"
        self.proj = Path(tempfile.mkdtemp(prefix="sdlc-proj-"))
        subprocess.run(["git", "init", "-q"], cwd=self.proj)
        self.env = {"HOME": str(self.home)}

    def tearDown(self):
        shutil.rmtree(self.home, ignore_errors=True)
        shutil.rmtree(self.proj, ignore_errors=True)

    def sdlc(self, *args, cwd=None):
        return subprocess.run([sys.executable, str(self.proj / "scripts" / "sdlc"), *args], cwd=cwd or self.proj,
                              capture_output=True, text=True, env=dict(os.environ, **self.env))

    def user_init(self):
        r = subprocess.run([sys.executable, str(self.user / "sdlc/bin/sdlc"), "init"], cwd=self.proj,
                           capture_output=True, text=True, env=dict(os.environ, **self.env))
        self.assertEqual(r.returncode, 0, r.stderr)

    def test_user_install_idempotent_and_uninstall(self):
        self.user.mkdir(parents=True)
        (self.user / "settings.json").write_text(json.dumps({"model": "x", "hooks": {"Stop": [{"hooks": []}]}}))
        for _ in range(2):
            r = install("--user", home=self.home)
            self.assertEqual(r.returncode, 0, r.stderr)
        for rel in ["sdlc/bin/sdlc", "sdlc/lib/sdlc_lib/gates.py", "sdlc/RULES.md", "hooks/sdlc/guard_write.py",
                    "skills/sdlc/SKILL.md", "skills/phase-intake/SKILL.md", "agents/reviewer.md"]:
            self.assertTrue((self.user / rel).exists(), rel)
        s = json.loads((self.user / "settings.json").read_text())
        self.assertEqual(s["model"], "x")
        self.assertIn("Stop", s["hooks"])
        cmds = [h["command"] for e in s["hooks"]["PreToolUse"] for h in e["hooks"]]
        self.assertEqual(len(cmds), 3)
        self.assertTrue(all(c.startswith('python3 "$HOME/.claude/hooks/sdlc/') for c in cmds))
        self.assertEqual(len(s["hooks"]["SessionStart"]), 1)
        r = install("--user", "--uninstall", home=self.home)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertFalse((self.user / "sdlc").exists())
        s = json.loads((self.user / "settings.json").read_text())
        self.assertEqual(s, {"model": "x", "hooks": {"Stop": [{"hooks": []}]}})

    def test_user_install_inactive_outside_sdlc_projects(self):
        install("--user", home=self.home)
        hook = self.user / "hooks/sdlc/guard_write.py"
        r = run_hook(hook, self.proj, {"file_path": str(self.proj / "docs/01-requirements/a.md")}, self.env)
        self.assertEqual(r.returncode, 0, r.stderr)  # SDLC プロジェクトでなければ何もしない
        r = run_hook(self.user / "hooks/sdlc/guard_bash.py", self.proj, {"command": "echo x > docs/a.md"}, self.env)
        self.assertEqual(r.returncode, 0, r.stderr)
        # ~/.claude の本体と設定ファイルは、どのプロジェクトからでも保護される
        r = run_hook(hook, self.proj, {"file_path": str(self.user / "sdlc/RULES.md")}, self.env)
        self.assertEqual(r.returncode, 2)
        r = run_hook(hook, self.proj, {"file_path": str(self.user / "settings.json")}, self.env)
        self.assertEqual(r.returncode, 2)
        r = run_hook(hook, self.proj, {"file_path": str(self.user / "skills/other/SKILL.md")}, self.env)
        self.assertEqual(r.returncode, 0)
        # SDLC でないプロジェクト自身の設定ファイルは制限しない
        r = run_hook(hook, self.proj, {"file_path": str(self.proj / ".claude/settings.json")}, self.env)
        self.assertEqual(r.returncode, 0)

    def test_user_install_project_init_and_gating(self):
        install("--user", home=self.home)
        self.user_init()
        for rel in ["scripts/sdlc", "scripts/approve", "state/phase.json", "docs/00-intake/.gitkeep"]:
            self.assertTrue((self.proj / rel).exists(), rel)
        self.assertIn("@~/.claude/sdlc/RULES.md", (self.proj / "CLAUDE.md").read_text())
        self.assertIn("workflow: sdlc", (self.proj / "project.yaml").read_text())
        self.assertIn(".sdlc/", (self.proj / ".gitignore").read_text())
        r = self.sdlc("where")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("配置: user", r.stdout)
        # project.yaml に workflow: sdlc があるので ~/.claude の hooks が工程を検査する
        hook = self.user / "hooks/sdlc/guard_write.py"
        r = run_hook(hook, self.proj, {"file_path": str(self.proj / "docs/01-requirements/a.md")}, self.env)
        self.assertEqual(r.returncode, 2)
        r = run_hook(hook, self.proj, {"file_path": str(self.proj / "scripts/approve")}, self.env)
        self.assertEqual(r.returncode, 2)

    def test_project_overrides_user(self):
        install("--user", home=self.home)
        r = install("--project", str(self.proj), home=self.home)
        self.assertEqual(r.returncode, 0, r.stderr)
        s = json.loads((self.proj / ".claude/settings.json").read_text())
        cmds = [h["command"] for e in s["hooks"]["PreToolUse"] for h in e["hooks"]]
        self.assertTrue(all("$CLAUDE_PROJECT_DIR/.claude/hooks/sdlc/" in c for c in cmds))
        self.user_init()
        self.assertIn("@.claude/sdlc/RULES.md", (self.proj / "CLAUDE.md").read_text())
        self.assertNotIn("@~/.claude/sdlc/RULES.md", (self.proj / "CLAUDE.md").read_text())
        target = {"file_path": str(self.proj / "docs/01-requirements/a.md")}
        # ~/.claude 側の hook はプロジェクト版に任せて何もしない
        self.assertEqual(run_hook(self.user / "hooks/sdlc/guard_write.py", self.proj, target, self.env).returncode, 0)
        # プロジェクト版の hook が検査する
        self.assertEqual(run_hook(self.proj / ".claude/hooks/sdlc/guard_write.py", self.proj, target, self.env).returncode, 2)
        # 本体のファイルはファイル単位でプロジェクトが優先
        r = self.sdlc("path", "templates/review.md")
        self.assertEqual(Path(r.stdout.strip()), self.proj / ".claude/sdlc/templates/review.md")
        self.assertIn("配置: project", self.sdlc("where").stdout)

    def test_partial_file_override(self):
        install("--user", home=self.home)
        self.user_init()
        override = self.proj / ".claude/sdlc/checklists/05-implementation.md"
        override.parent.mkdir(parents=True)
        override.write_text("# プロジェクト版\n")
        self.assertEqual(Path(self.sdlc("path", "checklists/05-implementation.md").stdout.strip()), override)
        self.assertEqual(Path(self.sdlc("path", "checklists/04-detailed-design.md").stdout.strip()),
                         self.user / "sdlc/checklists/04-detailed-design.md")
        self.assertIn("checklists/05-implementation.md", self.sdlc("where").stdout)

    def test_project_only_install(self):
        r = install("--project", str(self.proj), home=self.home)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertFalse((self.user / "sdlc").exists())
        r = subprocess.run([sys.executable, str(self.proj / ".claude/sdlc/bin/sdlc"), "init"], cwd=self.proj,
                           capture_output=True, text=True, env=dict(os.environ, **self.env))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("配置: project", self.sdlc("where").stdout)
        # プロジェクト配置では project.yaml の記入前から本体・工程成果物が保護される
        hook = self.proj / ".claude/hooks/sdlc/guard_write.py"
        self.assertEqual(run_hook(hook, self.proj, {"file_path": str(self.proj / ".claude/sdlc/RULES.md")},
                                  self.env).returncode, 2)
        self.assertEqual(run_hook(hook, self.proj, {"file_path": str(self.proj / "docs/00-intake/a.md")},
                                  self.env).returncode, 2)

    def test_launcher_without_install(self):
        install("--user", home=self.home)
        self.user_init()
        shutil.rmtree(self.user / "sdlc")
        r = self.sdlc("status")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("見つかりません", r.stderr)

    def test_reinstall_keeps_local_changes(self):
        install("--user", home=self.home)
        self.user_init()
        override = self.proj / ".claude/sdlc/templates/review.md"
        override.parent.mkdir(parents=True)
        override.write_text("# 社内様式\n")
        r = install("--project", str(self.proj), home=self.home)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(override.read_text(), "# 社内様式\n")       # インストーラが置いていないファイルは残す
        self.assertIn("templates/review.md", r.stdout)
        edited = self.proj / ".claude/sdlc/checklists/05-implementation.md"
        edited.write_text("# 編集\n")
        untouched = self.proj / ".claude/sdlc/phase-flow.md"
        untouched.write_text(untouched.read_text() + "x")            # 変更扱い
        untouched2 = self.proj / ".claude/sdlc/RULES.md"
        original = untouched2.read_text()
        r = install("--project", str(self.proj), home=self.home)
        self.assertEqual(edited.read_text(), "# 編集\n")
        self.assertEqual(untouched2.read_text(), original)
        r = install("--project", str(self.proj), "--force", home=self.home)
        self.assertNotEqual(edited.read_text(), "# 編集\n")
        r = install("--project", str(self.proj), "--uninstall", home=self.home)
        self.assertFalse((self.proj / ".claude/sdlc/bin").exists())
        self.assertFalse((self.proj / ".claude/hooks/sdlc").exists())

    def test_source_settings_match_installer(self):
        sys.path.insert(0, str(REPO))
        import importlib
        inst = importlib.import_module("install")
        expected = inst.merged_settings({}, "project")
        actual = json.loads((REPO / ".claude/settings.json").read_text(encoding="utf-8"))
        self.assertEqual(actual, expected)


if __name__ == "__main__":
    unittest.main()
