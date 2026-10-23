"""ワークフロー本体の自己テスト。  実行: python3 -m unittest discover -s scripts/tests -v"""
from __future__ import annotations

import contextlib
import importlib.machinery
import io
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import textwrap
import unittest
from argparse import Namespace
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "scripts"
HOOKS = REPO / ".claude" / "hooks"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "intake"))

from sdlc_lib import defects, fm, gates, state, tasks, tdd, trace  # noqa: E402
from sdlc_lib.config import load_project, phase_of_path, project_problems  # noqa: E402


def load_script(name):
    loader = importlib.machinery.SourceFileLoader(name, str(SCRIPTS / name))
    spec = importlib.util.spec_from_loader(name, loader)
    mod = importlib.util.module_from_spec(spec)
    loader.exec_module(mod)
    return mod


approve = load_script("approve")

PROJECT_YAML = """\
language: python
test_framework: unittest
commands:
  test: "python3 -m unittest discover -s tests -t . -q"
src_dir: src
test_dir: tests
"""


class Repo:
    """一時 git リポジトリ上のサンプルプロジェクト。"""

    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="sdlc-test-"))
        self.run("git", "init", "-q")
        self.run("git", "config", "user.email", "t@example.com")
        self.run("git", "config", "user.name", "tester")
        self.write("project.yaml", PROJECT_YAML)
        self.write(".gitignore", "__pycache__/\n.sdlc/\n")
        for d in ("src", "tests", "tests/unit", "tests/regression", "tests/integration", "tests/system"):
            self.write(f"{d}/__init__.py", "")
        self.write("input/req.md", "# 要求\n\n- 足し算ができること\n")
        state.save(self.dir, dict(state.DEFAULT_STATE))

    def run(self, *cmd, **kw):
        return subprocess.run(cmd, cwd=self.dir, capture_output=True, text=True, **kw)

    def write(self, rel, text):
        p = self.dir / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")

    @property
    def proj(self):
        return load_project(self.dir)

    def approve(self, phase, decision="approve"):
        a = Namespace(target=phase, defect_id=None, decision=decision, approver="tester", comment="ok", yes=True)
        with contextlib.redirect_stdout(io.StringIO()):
            approve.approve_phase(self.dir, self.proj, a)

    def hook(self, script, tool_input, env=None):
        data = {"tool_input": tool_input, "cwd": str(self.dir)}
        e = dict(os.environ, CLAUDE_PROJECT_DIR=str(self.dir))
        e.pop("SDLC_MAINTENANCE", None)
        e.update(env or {})
        return subprocess.run([sys.executable, str(HOOKS / script)], input=json.dumps(data), capture_output=True,
                              text=True, cwd=self.dir, env=e)

    def can_write(self, rel, **kw):
        r = self.hook("guard_write.py", {"file_path": str(self.dir / rel), "content": "x"}, **kw)
        return r.returncode == 0, r.stderr

    def cleanup(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class Base(unittest.TestCase):
    def setUp(self):
        self.r = Repo()
        self._cwd = os.getcwd()
        os.chdir(self.r.dir)

    def tearDown(self):
        os.chdir(self._cwd)
        self.r.cleanup()

    def approve_upto(self, k):
        """工程 0..k の成果物を用意して承認する。"""
        docs = {
            0: ("docs/00-intake/req.md", "### SRC-001\n\n> 足し算\n"),
            1: ("docs/01-requirements/requirements.md", "### REQ-001: 足し算\n- 追跡元: SRC-001\n"),
            2: ("docs/02-system-requirements/sr.md",
                "### SR-001: add\n- 追跡元: REQ-001\n\n### AT-001: add の受入\n- 追跡元: SR-001\n"),
            3: ("docs/03-basic-design/bd.md", "### BD-001: calc\n- 追跡元: SR-001\n"),
            4: ("docs/04-detailed-design/dd.md",
                "### DD-001: add()\n- 追跡元: BD-001\n\n### UT-001: add のテスト\n- 追跡元: DD-001\n"),
        }
        for i in range(k + 1):
            if i in docs:
                self.r.write(*docs[i])
            if i == 4:
                self.r.write("tests/unit/test_calc.py", textwrap.dedent("""\
                    # trace: UT-001
                    import unittest
                    from src.calc import add

                    class T(unittest.TestCase):
                        def test_add(self):
                            self.assertEqual(add(1, 2), 3)
                    """))
            self.r.approve(i)


class TestConfig(Base):
    def test_project_and_areas(self):
        p = self.r.proj
        self.assertEqual(project_problems(p), [])
        self.assertEqual(p["test_dirs"]["unit"], "tests/unit")
        self.assertEqual(phase_of_path("src/a.py", p), "05-implementation")
        self.assertEqual(phase_of_path("tests/unit/t.py", p), "04-detailed-design")
        self.assertEqual(phase_of_path("tests/regression/t.py", p), None)
        self.assertEqual(phase_of_path("input/x.xlsx", p), "00-intake")

    def test_reserved_dirs(self):
        self.r.write("project.yaml", PROJECT_YAML.replace("src_dir: src", "src_dir: docs/src"))
        self.assertTrue(any("予約" in x for x in project_problems(self.r.proj)))


class TestGates(Base):
    def test_hash_invalidation(self):
        self.approve_upto(0)
        self.assertTrue(gates.gate_status(self.r.dir, self.r.proj, 0)["valid"])
        self.r.write("docs/00-intake/req.md", "changed")
        st = gates.gate_status(self.r.dir, self.r.proj, 0)
        self.assertFalse(st["valid"])
        self.assertIn("ハッシュ不一致", st["reason"])

    def test_cannot_approve_without_prereq(self):
        self.r.write("docs/01-requirements/r.md", "x")
        with self.assertRaises(SystemExit):
            self.r.approve(1)

    def test_rejected_not_valid(self):
        self.r.write("docs/00-intake/req.md", "x")
        self.r.approve(0, "reject")
        self.assertFalse(gates.gate_status(self.r.dir, self.r.proj, 0)["valid"])

    def test_history_kept(self):
        self.approve_upto(0)
        self.r.write("docs/00-intake/req.md", "v2")
        self.r.approve(0)
        _, body = fm.read(gates.gate_path(self.r.dir, 0))
        self.assertIn("## 履歴", body)


class TestGuardWrite(Base):
    def test_gates_dir_denied(self):
        ok, msg = self.r.can_write("state/gates/00-intake.md")
        self.assertFalse(ok)
        ok, _ = self.r.can_write("state/gates/x.md", env={"SDLC_MAINTENANCE": "1"})
        self.assertFalse(ok)

    def test_protected_paths(self):
        self.assertFalse(self.r.can_write(".claude/settings.json")[0])
        self.assertFalse(self.r.can_write("scripts/sdlc_lib/gates.py")[0])
        self.assertTrue(self.r.can_write(".claude/settings.json", env={"SDLC_MAINTENANCE": "1"})[0])

    def test_phase_order(self):
        self.assertTrue(self.r.can_write("docs/00-intake/a.md")[0])
        ok, msg = self.r.can_write("docs/01-requirements/a.md")
        self.assertFalse(ok)
        self.assertIn("00-intake", msg)
        self.approve_upto(0)
        self.assertTrue(self.r.can_write("docs/01-requirements/a.md")[0])
        # 承認済み工程は変更不可
        ok, msg = self.r.can_write("docs/00-intake/a.md")
        self.assertFalse(ok)
        self.assertIn("承認済み", msg)

    def test_free_paths(self):
        self.assertTrue(self.r.can_write("state/tasks.yaml")[0])
        self.assertTrue(self.r.can_write("tests/regression/test_x.py")[0])
        self.assertTrue(self.r.can_write("README.md")[0])

    def test_tdd(self):
        self.approve_upto(4)  # 単体テストは Red（src/calc.py が無い）
        ok, msg = self.r.can_write("src/calc.py")
        self.assertTrue(ok, msg)
        self.r.write("src/calc.py", "def add(a, b):\n    return a + b\n")
        ok, msg = self.r.can_write("src/calc.py")
        self.assertTrue(ok, msg)  # Red 観測済み → Refactor 許可
        tdd.reset(self.r.dir)
        ok, msg = self.r.can_write("src/calc.py")
        self.assertFalse(ok)
        self.assertIn("TDD", msg)

    def test_unit_tests_frozen_after_detailed_design(self):
        self.approve_upto(4)
        self.assertFalse(self.r.can_write("tests/unit/test_calc.py")[0])


class TestGuardBashAndSkill(Base):
    def bash(self, cmd):
        return self.r.hook("guard_bash.py", {"command": cmd}).returncode == 0

    def test_bash(self):
        self.assertFalse(self.bash("cat state/gates/00-intake.md"))
        self.assertFalse(self.bash("scripts/approve 0 --yes"))
        self.assertFalse(self.bash("echo x > docs/01-requirements/a.md"))
        self.assertFalse(self.bash("sed -i s/a/b/ src/calc.py"))
        self.assertFalse(self.bash("cp /tmp/x .claude/hooks/guard_write.py"))
        self.assertTrue(self.bash("python3 -m unittest discover -s tests 2>&1 | tail"))
        self.assertTrue(self.bash("scripts/sdlc status"))
        self.assertTrue(self.bash("git merge --no-ff sdlc/T-001"))
        self.assertTrue(self.bash("ls docs > /dev/null"))

    def test_skill(self):
        r = self.r.hook("guard_skill.py", {"skill": "phase-requirements"})
        self.assertEqual(r.returncode, 2)
        self.approve_upto(0)
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "phase-requirements"}).returncode, 0)
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "sdlc"}).returncode, 0)


class TestTasks(unittest.TestCase):
    def data(self):
        return {"tasks": [
            {"id": "T1", "agent": "analyst", "files": ["docs/a.md"]},
            {"id": "T2", "agent": "analyst", "files": ["docs/a.md"]},
            {"id": "T3", "agent": "analyst", "files": ["docs/b.md"]},
            {"id": "T4", "agent": "analyst", "files": ["docs/c.md"], "depends_on": ["T1"]},
        ]}

    def test_ready_respects_files_and_deps(self):
        d = self.data()
        for t in d["tasks"]:
            t.setdefault("status", "pending")
            t.setdefault("depends_on", [])
        self.assertEqual([t["id"] for t in tasks.ready(d, 2)], ["T1", "T3"])
        tasks.set_status(d, "T1", "in_progress")
        self.assertEqual([t["id"] for t in tasks.ready(d, 2)], ["T3"])
        tasks.set_status(d, "T1", "done")
        self.assertEqual([t["id"] for t in tasks.ready(d, 2)], ["T2", "T3"])

    def test_cycle(self):
        d = {"tasks": [{"id": "A", "agent": "analyst", "files": ["x"], "depends_on": ["B"], "status": "pending"},
                       {"id": "B", "agent": "analyst", "files": ["y"], "depends_on": ["A"], "status": "pending"}]}
        self.assertTrue(any("循環" in e for e in tasks.validate(d)))

    def test_recover(self):
        d = {"tasks": [{"id": "A", "status": "in_progress"}]}
        self.assertEqual(tasks.recover(d), ["A"])
        self.assertEqual(d["tasks"][0]["status"], "pending")


class TestTrace(Base):
    def test_check_and_impact(self):
        self.approve_upto(4)
        g = trace.build(self.r.dir, self.r.proj)
        self.assertEqual(trace.check(g, 4), [])
        issues = trace.check(g, 5)
        self.assertTrue(any("DD-001" in i and "コード" in i for i in issues))
        imp = trace.impact(g, ["SR-001"])
        self.assertIn("UT-001", imp["ids"])
        self.assertIn("tests/unit/test_calc.py", imp["files"])
        self.assertIn("REQ-001", trace.upstream(g, ["DD-001"]))

    def test_exclude_and_manual(self):
        self.r.write("docs/00-intake/a.md", "### SRC-001\n\n### SRC-002\n")
        self.r.write("docs/01-requirements/r.md", "### REQ-001: x\n- 追跡元: SRC-001\n\n## 対象外\n- 除外: SRC-002（表題）\n")
        g = trace.build(self.r.dir, self.r.proj)
        self.assertEqual([i for i in trace.check(g, 1)], [])
        self.r.write("docs/02-system-requirements/at.md", "### AT-001: y\n- 追跡元: REQ-001\n")
        self.r.write("docs/07-system-test/manual-tests.md", "## 記録1\n- trace: AT-001\n- 結果: 合格\n")
        g = trace.build(self.r.dir, self.r.proj)
        self.assertFalse(any("AT-001" in i for i in trace.check(g, 7)))

    def test_dangling(self):
        self.r.write("docs/01-requirements/r.md", "### REQ-002: x\n- 追跡元: SRC-999\n")
        issues = trace.check(trace.build(self.r.dir, self.r.proj), 1)
        self.assertTrue(any("未定義" in i for i in issues))


class TestDefectFlow(Base):
    def sdlc(self, *args):
        r = subprocess.run([sys.executable, str(SCRIPTS / "sdlc"), *args], cwd=self.r.dir, capture_output=True, text=True,
                           env=dict(os.environ, CLAUDE_PROJECT_DIR=str(self.r.dir)))
        return r

    def test_full_defect_cycle(self):
        self.approve_upto(4)
        self.r.write("src/calc.py", "def add(a, b):\n    # trace: DD-001\n    return a - b\n")
        self.r.write("tests/unit/test_calc.py", (self.r.dir / "tests/unit/test_calc.py").read_text())
        st = state.load(self.r.dir)
        st["current_phase"] = "06-integration-test"
        state.save(self.r.dir, st)
        # 起票（再現テストは regression に）
        r = self.sdlc("defect", "new", "--title", "add が引き算", "--detected-in", "06", "--by", "auto",
                      "--test", "tests/regression/test_def_001.py")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.r.write("tests/regression/test_def_001.py",
                     "import unittest\nfrom src.calc import add\nclass T(unittest.TestCase):\n"
                     "    def test(self):\n        self.assertEqual(add(2, 2), 4)\n")
        meta, body = defects.load(self.r.dir, "DEF-001")
        meta.update(cause_ids=["DD-001"], injected_phase="05-implementation", escaped_phase="05-implementation")
        defects.save(self.r.dir, "DEF-001", meta, body)
        r = self.sdlc("defect", "impact", "DEF-001", "--write")
        self.assertEqual(r.returncode, 0, r.stderr)
        meta, _ = defects.load(self.r.dir, "DEF-001")
        self.assertIn("05-implementation", meta["rework_phases"])
        # 実装工程の承認（コードが存在する状態）
        self.r.approve(5)
        self.assertFalse(self.r.can_write("src/calc.py")[0])
        self.assertEqual(self.sdlc("defect", "analyzed", "DEF-001").returncode, 0)
        self.assertNotEqual(self.sdlc("defect", "start", "DEF-001").returncode, 0)  # 未承認
        with contextlib.redirect_stdout(io.StringIO()):
            approve.approve_defect(self.r.dir, self.r.proj, Namespace(
                target="defect", defect_id="DEF-001", decision="approve", approver="tester", comment="", yes=True))
        self.assertIn("再承認要", gates.gate_status(self.r.dir, self.r.proj, 5)["reason"])
        r = self.sdlc("defect", "start", "DEF-001")
        self.assertEqual(r.returncode, 0, r.stderr)
        tdd.reset(self.r.dir)
        ok, msg = self.r.can_write("src/calc.py")
        self.assertTrue(ok, msg)  # 再現テストが Red
        self.r.write("src/calc.py", "def add(a, b):\n    # trace: DD-001\n    return a + b\n")
        self.assertNotEqual(self.sdlc("defect", "close", "DEF-001").returncode, 0)  # 工程5未再承認
        self.r.approve(5)
        self.assertEqual(self.sdlc("phase", "advance").returncode, 0)
        r = self.sdlc("defect", "close", "DEF-001")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        st = state.load(self.r.dir)
        self.assertEqual((st["mode"], st["current_phase"]), ("normal", "06-integration-test"))

    def test_status_next_action(self):
        r = self.sdlc("status", "--json")
        self.assertEqual(json.loads(r.stdout)["next"]["action"], "run_phase")
        self.approve_upto(0)
        self.assertEqual(json.loads(self.sdlc("status", "--json").stdout)["next"]["action"], "advance")
        self.assertEqual(self.sdlc("phase", "advance").returncode, 0)
        self.sdlc("phase", "step", "gate")
        self.assertEqual(json.loads(self.sdlc("status", "--json").stdout)["next"]["action"], "await_approval")


class TestAdapters(Base):
    def convert(self, rel):
        import convert
        with contextlib.redirect_stdout(io.StringIO()):
            convert.main([rel])
        return (self.r.dir / "docs/00-intake" / (Path(rel).stem + ".md")).read_text(encoding="utf-8")

    def test_md(self):
        self.r.write("input/spec.md", "# 概要\n\n## 機能\n\n- ログインできる\n- 検索できる\n\n| ID | 内容 |\n|---|---|\n| 1 | 帳票 |\n")
        out = self.convert("input/spec.md")
        self.assertIn("SRC-001", out)
        self.assertIn("概要 > 機能 / L5", out)
        self.assertIn("内容: 帳票", out)

    def test_xlsx_merged(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "要求一覧"
        ws.append(["分類", "ID", "要求"])
        ws.append(["認証", "R1", "ログインできる"])
        ws.append([None, "R2", "ログアウトできる"])
        ws.merge_cells("A2:A3")
        wb.save(self.r.dir / "input/req.xlsx")
        out = self.convert("input/req.xlsx")
        self.assertIn("要求一覧!A3:C3", out)
        self.assertIn("分類: 認証\nID: R2", out.replace("> ", ""))
        self.assertIn("確信度: 中", out)

    def test_docx(self):
        import docx
        d = docx.Document()
        d.add_heading("機能要求", 1)
        d.add_paragraph("利用者は検索できること。")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "項目", "内容"
        t.cell(1, 0).text, t.cell(1, 1).text = "応答", "3秒以内"
        d.save(self.r.dir / "input/req.docx")
        out = self.convert("input/req.docx")
        self.assertIn("機能要求 / 段落2", out)
        self.assertIn("表1 行2", out)
        self.assertIn("内容: 3秒以内", out)

    def test_pptx(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "権限"
        rows = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        rows.cell(0, 0).text, rows.cell(0, 1).text = "機能", "管理者"
        rows.cell(1, 0).text, rows.cell(1, 1).text = "削除", "○"
        prs.save(self.r.dir / "input/deck.pptx")
        out = self.convert("input/deck.pptx")
        self.assertIn("スライド1「権限」", out)
        self.assertIn("管理者: ○", out)
        self.assertIn("確信度: 中", out)

    def test_refuse_after_gate(self):
        self.approve_upto(0)
        with self.assertRaises(SystemExit):
            self.convert("input/req.md")


if __name__ == "__main__":
    unittest.main()
